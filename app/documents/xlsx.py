"""تصدير Excel (§٨) — يُعيد استخدام **نفس** دوال بناء السياق التي تُغذّي
مستندات PDF (`reports_doc.py` · `statement.py`): لا إعادة اشتقاق أي منطق
محاسبي أو تنسيق أرقام هنا، فقط تحويل السياق الجاهز إلى ورقة عمل.

`report_to_xlsx` يعمل مع أي سياق `report_table.html` (أعمدة/صفوف/ملخّص).
`statement_to_xlsx` يعمل مع سياق كشف الحساب (رصيد افتتاحي/ختامي + رصيد
متحرك). كلاهما يُخرج bytes جاهزة للإرسال بـ`media_type` الصحيح.

القيم في السياق **نصوص مُنسَّقة مسبقاً** (`format_amount`) — تُكتب كما هي؛
الاتساق مع PDF أهم من كون الخلية رقمية في Excel (نفس مبدأ "مصدر واحد").
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14)
_RTL = Alignment(horizontal="right", vertical="center", wrap_text=True)


def _new_sheet(title: str):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "تقرير"
    ws.sheet_view.rightToLeft = True
    return wb, ws


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 3, 12), 60)


def _finalize(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_to_xlsx(ctx: dict) -> bytes:
    """سياق `report_table.html`: `doc_title` · `doc_subtitle` · `columns` ·
    `rows` (قوائم نصوص) · `summary` (قائمة {label, lines:[{k,v}]})."""
    wb, ws = _new_sheet(ctx.get("doc_title", "تقرير"))

    ws.append([ctx.get("doc_title", "تقرير")])
    ws["A1"].font = _TITLE_FONT
    if ctx.get("doc_subtitle"):
        ws.append([ctx["doc_subtitle"]])
    ws.append([f"{ctx.get('company_name', 'ومضة')} — أُصدر {ctx.get('generated_at', '')}"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(list(ctx.get("columns", [])))
    for cell in ws[header_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _RTL

    for row in ctx.get("rows", []):
        ws.append(list(row))

    for item in ctx.get("summary") or []:
        ws.append([])
        parts = "، ".join(f"{ln['k']}: {ln['v']}" for ln in item.get("lines", []))
        ws.append([item.get("label", ""), parts])
        ws[ws.max_row][0].font = Font(bold=True)

    for row in ws.iter_rows():
        for cell in row:
            if cell.alignment.horizontal is None:
                cell.alignment = _RTL
    _autosize(ws)
    return _finalize(wb)


def statement_to_xlsx(ctx: dict) -> bytes:
    """سياق كشف الحساب: `doc_title` · `doc_subtitle` · `opening_balance` ·
    `closing_balance` · `rows` (قوائم dict: entry_no/entry_date/operation_type/
    description/debit/credit/running)."""
    wb, ws = _new_sheet(ctx.get("doc_title", "كشف حساب"))

    ws.append([ctx.get("doc_title", "كشف حساب")])
    ws["A1"].font = _TITLE_FONT
    if ctx.get("doc_subtitle"):
        ws.append([ctx["doc_subtitle"]])
    ws.append([f"{ctx.get('company_name', 'ومضة')} — أُصدر {ctx.get('generated_at', '')}"])
    ws.append([])
    ws.append(["الرصيد الافتتاحي", ctx.get("opening_balance", "")])
    ws.append(["الرصيد الختامي", ctx.get("closing_balance", "")])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(["رقم القيد", "التاريخ", "النوع", "البيان", "مدين", "دائن", "الرصيد المتحرك"])
    for cell in ws[header_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _RTL

    for r in ctx.get("rows", []):
        ws.append([
            r.get("entry_no"), r.get("entry_date"), r.get("operation_type"), r.get("description"),
            r.get("debit"), r.get("credit"), r.get("running"),
        ])

    for row in ws.iter_rows():
        for cell in row:
            if cell.alignment.horizontal is None:
                cell.alignment = _RTL
    _autosize(ws)
    return _finalize(wb)
